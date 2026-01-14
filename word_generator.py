"""
Модуль для генерации Word документов
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import json
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional
import io
import re


GOST_FONT_NAME = "Times New Roman"
GOST_FONT_SIZE = 14
GOST_LINE_SPACING = 1.5
GOST_FIRST_LINE_INDENT_CM = 1.25


class SafeDict(dict):
    """Словарь для безопасной подстановки значений в шаблон."""

    def __missing__(self, key: str) -> str:
        return ""


def normalize_date(value: str) -> str:
    """Преобразование даты в формат ДД.ММ.ГГГГ при возможности."""
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", value)
    if match:
        year, month, day = match.groups()
        return f"{day}.{month}.{year}"
    return value


def render_template_text(text: str, fields_data: Dict[str, Any]) -> str:
    """Подстановка значений в текстовый шаблон."""
    prepared = {}
    for key, value in fields_data.items():
        if value is None:
            prepared[key] = ""
        elif isinstance(value, str):
            prepared[key] = normalize_date(value)
        else:
            prepared[key] = str(value)
    return text.format_map(SafeDict(prepared)).strip()


def apply_gost_styles(document: Document) -> None:
    """Настройка полей и базового стиля по ГОСТ."""
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    style = document.styles["Normal"]
    style.font.name = GOST_FONT_NAME
    style.font.size = Pt(GOST_FONT_SIZE)


def add_gost_paragraph(document: Document, text: str) -> None:
    """Добавить абзац с ГОСТ-оформлением."""
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(GOST_FIRST_LINE_INDENT_CM)
    paragraph.paragraph_format.line_spacing = GOST_LINE_SPACING


def add_heading_paragraph(document: Document, text: str, center: bool = False) -> None:
    """Добавить заголовок по ГОСТ."""
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = GOST_LINE_SPACING
    for run in paragraph.runs:
        run.bold = True


def add_list_items(document: Document, items: Iterable[str]) -> None:
    """Добавить нумерованный список."""
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph(f"{index}. {item}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0.5)
        paragraph.paragraph_format.line_spacing = GOST_LINE_SPACING


def add_signature_block(
    document: Document,
    left_label: str,
    right_label: str,
    left_name: str,
    right_name: str,
) -> None:
    """Добавить блок подписей сторон."""
    table = document.add_table(rows=2, cols=2)
    table.autofit = True
    table.cell(0, 0).text = left_label
    table.cell(0, 1).text = right_label
    table.cell(1, 0).text = f"__________________ {left_name}".strip()
    table.cell(1, 1).text = f"__________________ {right_name}".strip()


def build_document_from_structure(
    document: Document,
    structure: List[Dict[str, Any]],
    fields_data: Dict[str, Any],
) -> None:
    """Сборка документа по структуре, сохраненной в БД."""
    for block in structure:
        block_type = block.get("type")
        text = block.get("text", "")
        rendered = render_template_text(text, fields_data) if text else ""

        if block_type == "title":
            add_heading_paragraph(document, rendered.upper(), center=True)
        elif block_type == "heading":
            add_heading_paragraph(document, rendered, center=False)
        elif block_type == "paragraph":
            if rendered:
                add_gost_paragraph(document, rendered)
        elif block_type == "list":
            items = [
                render_template_text(item, fields_data)
                for item in block.get("items", [])
            ]
            add_list_items(document, [item for item in items if item])
        elif block_type == "signature":
            left_label = render_template_text(block.get("left_label", ""), fields_data)
            right_label = render_template_text(block.get("right_label", ""), fields_data)
            left_name = render_template_text(block.get("left_name", ""), fields_data)
            right_name = render_template_text(block.get("right_name", ""), fields_data)
            add_signature_block(
                document,
                left_label=left_label,
                right_label=right_label,
                left_name=left_name,
                right_name=right_name,
            )
        elif block_type == "spacer":
            document.add_paragraph()


def create_fallback_document(document_name: str, fields_data: Dict[str, Any]) -> bytes:
    """Создание простого документа на случай ошибки."""
    doc = Document()
    apply_gost_styles(doc)
    add_heading_paragraph(doc, document_name, center=True)
    add_gost_paragraph(
        doc,
        f"Документ создан автоматически {datetime.now().strftime('%d.%m.%Y %H:%M')}",
    )
    for key, value in fields_data.items():
        if value:
            add_gost_paragraph(doc, f"{key}: {value}")
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


def generate_word_document(
    template_data: Dict[str, Any],
    fields_data: Dict[str, Any],
    document_name: str,
    user_info: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Функция для генерации Word документа по структуре из БД."""
    try:
        document = Document()
        apply_gost_styles(document)

        structure_raw = template_data.get("content_json")
        content_text = template_data.get("content_text")
        if structure_raw:
            structure = json.loads(structure_raw)
        elif content_text:
            structure = [
                {"type": "paragraph", "text": line}
                for line in content_text.splitlines()
            ]
        else:
            structure = [
                {"type": "title", "text": document_name},
                {"type": "paragraph", "text": "Шаблон документа не содержит структуры."},
            ]

        if not any(block.get("type") == "title" for block in structure):
            structure.insert(0, {"type": "title", "text": document_name})

        build_document_from_structure(document, structure, fields_data)

        output = io.BytesIO()
        document.save(output)
        output.seek(0)
        return output.getvalue()
    except Exception as exc:  # pragma: no cover - fallback
        print(f"Ошибка генерации Word документа: {exc}")
        return create_fallback_document(document_name, fields_data)